"""proposal@1: numbered scope/deliverables/timeline, one priced table, validity term."""

import json
from pathlib import Path

import pytest
from document_production.cli import main
from document_production.jobs import run_proposal_job
from document_production.proposal import (
    build_proposal_ir,
    load_proposal_request_schema,
    proposal_request_schema,
)
from docx import Document
from pydantic import ValidationError

PROPOSAL = {
    "schema_version": "proposal-request.v1",
    "document_id": "pp-2026-001",
    "title": "ข้อเสนองานผลิตสื่อ",
    "locale": "th-TH",
    "currency": "THB",
    "customer_name": "ร้านตัวอย่าง",
    "issue_date": "2026-09-30",
    "valid_until": "2026-10-30",
    "objective": ["เพิ่มยอดขายช่วงเปิดตัวเสื้อลายใหม่"],
    "scope": ["ออกแบบภาพโปรโมต 30 ชิ้น", "ภาพสินค้า 3D 3 มุม"],
    "deliverables": [{"name": "ภาพโปรโมต", "description": "PNG + PDF สำหรับพิมพ์"}, {"name": "ไฟล์ต้นฉบับ"}],
    "timeline": [{"phase": "ออกแบบ", "duration": "5 วันทำการ"}, {"phase": "แก้ไข", "duration": "2 รอบ"}],
    "line_items": [
        {"description": "ภาพโปรโมต 30 ชิ้น", "quantity": "1", "unit_price": "4500.00", "tax_rate": "7"},
        {"description": "ภาพ 3D", "quantity": "3", "unit_price": "800.00", "tax_rate": "7"},
    ],
    "terms": ["ชำระ 50% ก่อนเริ่มงาน"],
}


def test_proposal_ir_sections_and_totals() -> None:
    ir = build_proposal_ir(PROPOSAL)
    assert [section.heading for section in ir.sections] == [
        "วัตถุประสงค์", "ขอบเขตงาน", "สิ่งที่ส่งมอบ", "ระยะเวลาดำเนินงาน", "ค่าใช้จ่าย", "เงื่อนไข",
    ]
    assert ir.sections[2].paragraphs == ["1. ภาพโปรโมต — PNG + PDF สำหรับพิมพ์", "2. ไฟล์ต้นฉบับ"]
    assert ir.sections[-1].paragraphs[-1] == "2. ข้อเสนอนี้มีผลถึงวันที่ 2026-10-30"
    assert sum(len(section.tables) for section in ir.sections) == 1
    assert ir.derived_totals == {"subtotal": "6900.00", "tax_total": "483.00", "grand_total": "7383.00"}


def test_validity_before_issue_date_is_rejected() -> None:
    with pytest.raises(ValidationError):
        build_proposal_ir({**PROPOSAL, "valid_until": "2026-09-01"})


def test_job_passes_totals_qa(tmp_path: Path) -> None:
    result = run_proposal_job(PROPOSAL, tmp_path, "pp")
    assert result.status == "accepted", result.qa
    text = " ".join(p.text for p in Document(str(Path(result.manifest_path).parent / "document.docx")).paragraphs)
    assert "ขอบเขตงาน" in text and "grand_total: 7383.00 THB" in text


def test_published_schema_matches_model() -> None:
    assert load_proposal_request_schema() == proposal_request_schema()


def test_cli_proposal(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    request = tmp_path / "proposal.json"
    request.write_text(json.dumps(PROPOSAL, ensure_ascii=False), encoding="utf-8")
    assert main(["proposal", "--input", str(request), "--output", str(tmp_path / "out"), "--idempotency-key", "c"]) == 0
    assert json.loads(capsys.readouterr().out)["status"] == "accepted"
