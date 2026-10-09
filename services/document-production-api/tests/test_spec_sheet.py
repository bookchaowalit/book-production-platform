"""spec-sheet@1: specification table without totals, verified preview images, Thai labels."""

import hashlib
import io
import json
from pathlib import Path

import pytest
from docx import Document
from PIL import Image
from pydantic import ValidationError
from pypdf import PdfReader

from document_production.cli import main
from document_production.jobs import run_spec_sheet_job
from document_production.recipe import RequestValidationError
from document_production.spec_sheet import (
    build_spec_sheet_ir,
    load_spec_sheet_request_schema,
    spec_sheet_request_schema,
)

REQUEST = {
    "schema_version": "spec-sheet-request.v1",
    "document_id": "spec-plate-001",
    "title": "Spec sheet แผ่นยึด",
    "locale": "th-TH",
    "customer_name": "ลูกค้าตัวอย่าง",
    "issue_date": "2026-09-30",
    "subject": "Mounting plate 137 x 83 x 7.5 mm",
    "specs": [
        {"label": "ความยาว", "value": "137.0", "unit": "mm"},
        {"label": "ปริมาตร", "value": "83774.54", "unit": "mm3"},
        {"label": "จำนวนรู", "value": "4"},
    ],
    "notes": ["เอกสารนี้ไม่ใช่การรับรองทางวิศวกรรม"],
}


def _png(path: Path) -> dict:
    buffer = io.BytesIO()
    Image.new("RGB", (320, 240), (200, 200, 200)).save(buffer, format="PNG")
    data = buffer.getvalue()
    path.write_bytes(data)
    return {"name": path.name, "media_type": "image/png", "content_hash": hashlib.sha256(data).hexdigest(),
            "byte_size": len(data)}


def test_ir_has_spec_table_notes_and_no_totals() -> None:
    ir = build_spec_sheet_ir(REQUEST)
    assert ir.template_ref == "spec-sheet@1"
    assert ir.derived_totals == {}
    table = ir.sections[1].tables[0]
    assert table.headers == ["รายการ", "ค่า", "หน่วย"]
    assert [row.cells for row in table.rows][2] == ["จำนวนรู", "4", "-"]
    assert ir.sections[-1].paragraphs == ["เอกสารนี้ไม่ใช่การรับรองทางวิศวกรรม"]


def test_english_locale_uses_english_headings() -> None:
    ir = build_spec_sheet_ir({**REQUEST, "locale": "en-US"})
    assert ir.sections[1].heading == "Specifications"


def test_job_embeds_verified_preview_and_uses_thai_labels(tmp_path: Path) -> None:
    ref = _png(tmp_path / "preview.png")
    result = run_spec_sheet_job({**REQUEST, "images": [ref]}, tmp_path / "out", "spec-1", image_root=tmp_path)
    assert result.status == "accepted", result.qa
    job_dir = Path(result.manifest_path).parent
    document = Document(str(job_dir / "document.docx"))
    assert len(document.inline_shapes) == 1
    text = " ".join(paragraph.text for paragraph in document.paragraphs)
    assert "ลูกค้า: ลูกค้าตัวอย่าง" in text
    assert "Customer:" not in text
    assert "Totals" not in text
    assert sum(len(page.images) for page in PdfReader(str(job_dir / "document.pdf")).pages) == 1


def test_tampered_preview_is_rejected(tmp_path: Path) -> None:
    ref = _png(tmp_path / "preview.png")
    (tmp_path / "preview.png").write_bytes(b"not the same image")
    with pytest.raises(RequestValidationError) as caught:
        build_spec_sheet_ir({**REQUEST, "images": [ref]}, image_root=tmp_path)
    assert caught.value.issues[0].code == "image_file_mismatch"


def test_request_validation() -> None:
    with pytest.raises(ValidationError):
        build_spec_sheet_ir({**REQUEST, "specs": []})
    with pytest.raises(ValidationError):
        build_spec_sheet_ir({**REQUEST, "unexpected": True})


def test_published_schema_matches_model() -> None:
    assert load_spec_sheet_request_schema() == spec_sheet_request_schema()


def test_cli_spec_sheet(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    request = tmp_path / "request.json"
    request.write_text(json.dumps(REQUEST, ensure_ascii=False), encoding="utf-8")
    assert main(["spec-sheet", "--input", str(request), "--output", str(tmp_path / "out"), "--idempotency-key", "c1"]) == 0
    assert json.loads(capsys.readouterr().out)["status"] == "accepted"
