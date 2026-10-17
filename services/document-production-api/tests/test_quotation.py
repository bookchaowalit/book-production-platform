"""quotation@1 reuses the invoice arithmetic with Thai table labels for th-* locales."""

import json
from pathlib import Path

import pytest
from document_production.cli import main
from document_production.recipe import RequestValidationError, build_document_ir
from docx import Document

QUOTATION = {
    "schema_version": "document-request.v1",
    "template_id": "quotation",
    "template_version": "1",
    "document_id": "qt-2026-001",
    "locale": "th-TH",
    "currency": "THB",
    "title": "ใบเสนอราคา",
    "customer_name": "ร้านตัวอย่าง",
    "issue_date": "2026-09-30",
    "line_items": [
        {"description": "ออกแบบภาพโปรโมต 30 ชิ้น", "quantity": "1", "unit_price": "4500.00", "tax_rate": "7"},
        {"description": "ภาพ 3D สินค้า", "quantity": "3", "unit_price": "800.00", "tax_rate": "7"},
    ],
    "notes": "ราคานี้ยืนราคา 30 วันนับจากวันที่ออกเอกสาร",
}


def test_quotation_totals_and_thai_labels() -> None:
    ir = build_document_ir(QUOTATION)
    assert ir.template_ref == "quotation@1"
    assert ir.assumptions["recipe"] == "quotation-v1"
    assert ir.derived_totals == {"subtotal": "6900.00", "tax_total": "483.00", "grand_total": "7383.00"}
    section = ir.sections[0]
    assert section.heading == "รายการเสนอราคา"
    assert section.tables[0].headers == ["รายการ", "จำนวน", "ราคาต่อหน่วย", "ภาษี %", "จำนวนเงิน"]
    assert section.paragraphs == ["ราคานี้ยืนราคา 30 วันนับจากวันที่ออกเอกสาร"]


def test_invoice_labels_are_unchanged() -> None:
    ir = build_document_ir({**QUOTATION, "template_id": "invoice"})
    assert ir.sections[0].heading == "Items"
    assert ir.assumptions["recipe"] == "invoice-v1"


def test_unknown_template_is_rejected() -> None:
    with pytest.raises(RequestValidationError):
        build_document_ir({**QUOTATION, "template_id": "receipt"})


def test_cli_renders_quotation(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    request = tmp_path / "quotation.json"
    request.write_text(json.dumps(QUOTATION, ensure_ascii=False), encoding="utf-8")
    assert main(["render", "--input", str(request), "--output", str(tmp_path / "out"), "--idempotency-key", "qt"]) == 0
    result = json.loads(capsys.readouterr().out)
    assert result["status"] == "accepted"
    job_dir = next((tmp_path / "out" / "jobs").iterdir())
    manifest = json.loads((job_dir / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["recipe"] == "quotation-v1"
    text = " ".join(paragraph.text for paragraph in Document(str(job_dir / "document.docx")).paragraphs)
    assert "ลูกค้า: ร้านตัวอย่าง" in text and "grand_total: 7383.00 THB" in text
