"""Local DOCX and PDF renderers for normalized document IR."""

from __future__ import annotations

import hashlib
import os
import tempfile
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from document_ir import ArtifactRef, DocumentImage, DocumentIR
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

SUPPORTED_TEMPLATES = frozenset({"invoice@1", "quotation@1", "proposal@1", "transcript@1", "report@1", "spec-sheet@1"})
PDF_FONT_NAME = "DocumentProductionUnicode"


class RenderError(RuntimeError):
    """Raised when a deterministic artifact cannot be rendered."""


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _header_lines(ir: DocumentIR) -> tuple[str, str]:
    """Customer and issue-date lines in the document language (Thai for th-* locales)."""

    if ir.locale.lower().startswith("th"):
        return f"ลูกค้า: {ir.customer_name}", f"วันที่ออกเอกสาร: {ir.issue_date.isoformat()}"
    return f"Customer: {ir.customer_name}", f"Issue date: {ir.issue_date.isoformat()}"


def _pdf_column_widths(ir: DocumentIR, columns: int) -> list[float]:
    if ir.template_ref == "spec-sheet@1" and columns == 3:
        return [80 * mm, 70 * mm, 30 * mm]
    return [65 * mm, 20 * mm, 28 * mm, 20 * mm, 28 * mm]


def _verified_image(image: DocumentImage) -> str:
    """Re-hash an embedded image right before rendering so a swapped file is never used."""

    path = Path(image.path)
    if not path.is_file() or _sha256(path) != image.content_hash:
        raise RenderError(f"embedded image changed or is missing: {path.name}")
    return str(path)


def _atomic_path(output_dir: Path, filename: str) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    final_path = output_dir / filename
    with tempfile.NamedTemporaryFile(prefix=f".{filename}.", suffix=".tmp", dir=output_dir, delete=False) as handle:
        return Path(handle.name), final_path


def _artifact(path: Path, filename: str, media_type: str, page_count: int | None = None) -> ArtifactRef:
    return ArtifactRef(
        filename=filename,
        media_type=media_type,
        content_hash=_sha256(path),
        byte_size=path.stat().st_size,
        page_count=page_count,
    )


def _footer_label(ir: DocumentIR) -> str:
    """Footer prefix shown before the page number on every page (DOCX and PDF)."""

    return f"{ir.document_id} · หน้า " if ir.locale.lower().startswith("th") else f"{ir.document_id} · Page "


def _add_docx_page_footer(document: Any, ir: DocumentIR) -> None:
    """Add ``<document_id> · Page <PAGE field>`` to the DOCX footer (Insert ▸ Page Number in Word)."""

    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run(_footer_label(ir))
    run = paragraph.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag is None:
            instruction = OxmlElement("w:instrText")
            instruction.set(qn("xml:space"), "preserve")
            instruction.text = text
            run._r.append(instruction)
        else:
            field = OxmlElement("w:fldChar")
            field.set(qn("w:fldCharType"), tag)
            run._r.append(field)


def _render_docx(ir: DocumentIR, output_dir: Path) -> ArtifactRef:
    temporary, final = _atomic_path(output_dir, "document.docx")
    try:
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        _add_docx_page_footer(document, ir)
        document.add_heading(ir.title, level=0)
        for line in _header_lines(ir):
            document.add_paragraph(line)
        for item_section in ir.sections:
            document.add_heading(item_section.heading, level=1)
            for paragraph in item_section.paragraphs:
                document.add_paragraph(paragraph)
            for item_table in item_section.tables:
                table = document.add_table(rows=1, cols=len(item_table.headers))
                table.style = "Table Grid"
                for cell, value in zip(table.rows[0].cells, item_table.headers):
                    cell.text = value
                for row in item_table.rows:
                    cells = table.add_row().cells
                    for cell, value in zip(cells, row.cells):
                        cell.text = value
            for image in item_section.images:
                document.add_picture(_verified_image(image), width=Inches(6.3))
                if image.caption:
                    document.add_paragraph(image.caption)
        # The Totals block is invoice-specific; templates without monetary
        # totals (e.g. transcript@1) publish an empty derived_totals mapping.
        if ir.derived_totals:
            document.add_heading("Totals", level=1)
            for name, value in ir.derived_totals.items():
                document.add_paragraph(f"{name}: {value} {ir.currency}")
        document.save(str(temporary))
        os.replace(temporary, final)
    except Exception as exc:
        temporary.unlink(missing_ok=True)
        raise RenderError(f"DOCX render failed: {exc}") from exc
    return _artifact(final, "document.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")


def _font_name() -> str:
    if PDF_FONT_NAME in pdfmetrics.getRegisteredFontNames():
        return PDF_FONT_NAME
    candidates = [
        os.environ.get("DOCUMENT_PRODUCTION_FONT"),
        "C:/Windows/Fonts/LeelawUI.ttf",
        "/usr/share/fonts/truetype/noto/NotoSansThai-Regular.ttf",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            pdfmetrics.registerFont(TTFont(PDF_FONT_NAME, candidate))
            return PDF_FONT_NAME
    return "Helvetica"


def _pdf_value(value: str, font_name: str) -> str:
    if font_name != "Helvetica":
        return escape(value)
    if any(ord(character) > 255 for character in value):
        raise RenderError("PDF font cannot encode non-Latin text; configure DOCUMENT_PRODUCTION_FONT")
    return escape(value)


def _render_pdf(ir: DocumentIR, output_dir: Path) -> ArtifactRef:
    temporary, final = _atomic_path(output_dir, "document.pdf")
    try:
        font_name = _font_name()
        styles = getSampleStyleSheet()
        body = ParagraphStyle("DocumentBody", parent=styles["BodyText"], fontName=font_name, fontSize=10, leading=14, alignment=TA_LEFT)
        heading = ParagraphStyle("DocumentHeading", parent=styles["Heading1"], fontName=font_name, fontSize=14, leading=18)
        story = [Paragraph(_pdf_value(ir.title, font_name), heading), Spacer(1, 4 * mm)]
        for line in _header_lines(ir):
            story.append(Paragraph(_pdf_value(line, font_name), body))
        for item_section in ir.sections:
            story.append(Spacer(1, 4 * mm))
            story.append(Paragraph(_pdf_value(item_section.heading, font_name), heading))
            for paragraph in item_section.paragraphs:
                story.append(Paragraph(_pdf_value(paragraph, font_name), body))
            for item_table in item_section.tables:
                data = [[Paragraph(_pdf_value(value, font_name), body) for value in item_table.headers]]
                data.extend([[Paragraph(_pdf_value(value, font_name), body) for value in row.cells] for row in item_table.rows])
                table = Table(data, repeatRows=1, colWidths=_pdf_column_widths(ir, len(item_table.headers)))
                table.setStyle(TableStyle([
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]))
                story.append(Spacer(1, 2 * mm))
                story.append(table)
            for image in item_section.images:
                story.append(Spacer(1, 3 * mm))
                story.append(Image(_verified_image(image), width=170 * mm, height=120 * mm, kind="proportional"))
                if image.caption:
                    story.append(Paragraph(_pdf_value(image.caption, font_name), body))
        if ir.derived_totals:
            story.append(Spacer(1, 4 * mm))
            story.append(Paragraph(_pdf_value("Totals", font_name), heading))
            for name, value in ir.derived_totals.items():
                story.append(Paragraph(_pdf_value(f"{name}: {value} {ir.currency}", font_name), body))

        document = SimpleDocTemplate(str(temporary), pagesize=A4, rightMargin=15 * mm, leftMargin=15 * mm, topMargin=15 * mm, bottomMargin=15 * mm)
        label = _pdf_value(_footer_label(ir), font_name)

        def footer(canvas: Any, template: Any) -> None:
            canvas.saveState()
            canvas.setFont(font_name, 8)
            canvas.drawRightString(A4[0] - 15 * mm, 8 * mm, f"{label}{template.page}")
            canvas.restoreState()

        document.build(story, onFirstPage=footer, onLaterPages=footer)
        os.replace(temporary, final)
    except Exception as exc:
        temporary.unlink(missing_ok=True)
        raise RenderError(f"PDF render failed: {exc}") from exc
    page_count = len(PdfReader(str(final)).pages)
    return _artifact(final, "document.pdf", "application/pdf", page_count=page_count)


def render_document(ir: DocumentIR, output_dir: str | Path) -> list[ArtifactRef]:
    """Render the supported IR into local DOCX and PDF artifacts."""

    if ir.template_ref not in SUPPORTED_TEMPLATES:
        raise RenderError(f"template {ir.template_ref!r} is not supported")
    output_path = Path(output_dir)
    return [_render_docx(ir, output_path), _render_pdf(ir, output_path)]
